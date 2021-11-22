from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
gradesheet_workbook = load_workbook("MAT110 FALL 2021 Gradesheet.xlsx")
mid_wb = load_workbook("MAT110 Midterm Exam Fall21 (Responses).xlsx")["Form Responses 1"]
quiz_wb = load_workbook("Quiz Marks.xlsx")["Quiz Marks"]
count=0
output = ""
information = Workbook()
forbot = Workbook()

bot = forbot.create_sheet("Quiz 1")
bot["A"+str(1)] = "Student ID"
bot["B"+str(1)] = "Name"
bot["C"+str(1)] = "G-Suite"
bot["D"+str(1)] = "BUX Username"
bot["E"+str(1)] = "Quiz Mark"
botcount = 1

for gradesheet_ws in gradesheet_workbook.worksheets:
    infos = information.create_sheet(gradesheet_ws.title)
    infos["A"+str(1)] = "Student ID"
    infos["B"+str(1)] = "Name"
    infos["C"+str(1)] = "G-Suite"
    infos["D"+str(1)] = "BUX Username"

    count_sec=0
    loss_count_sec=0

    absent=[]
    for row in range(4, 60):

        got = None
        id_char  = get_column_letter(2)+str(row)
        name_char  = get_column_letter(3)+str(row)
        id = str(gradesheet_ws[id_char].value)
        name = str(gradesheet_ws[name_char].value)
        if id!="None":
            got=False
        if id=="None":
            break
        botcount+=1
        infos["A"+str(row-2)] = id
        infos["B"+str(row-2)] = name
        bot["A"+str(botcount)] = id
        bot["B"+str(botcount)] = name
        for mid_row in range(2,845):
            mid_id_char = "F"+str(mid_row)
            mid_id = str(mid_wb[mid_id_char].value)
            if mid_id in id:
                got = True
                count_sec+=1
                mid_email_char = "B"+str(mid_row)
                mid_email2_char = get_column_letter(3)+str(mid_row)
                mid_bux_char = get_column_letter(4)+str(mid_row)
                mid_email = str(mid_wb[mid_email_char].value)
                for quiz_row in range(2, 1107):
                    quiz_bux_char = "C"+str(quiz_row)
                    quiz_email_char = get_column_letter(2)+str(quiz_row)
                    quiz_marks_char = "AL"+str(quiz_row)
                    quiz_email = str(quiz_wb[quiz_email_char].value)
                    quiz_marks = quiz_wb[quiz_marks_char].value
                    quiz_bux = quiz_wb[quiz_bux_char].value
                    if quiz_marks=="Not Attempted":
                        quiz_marks=0
                    else: quiz_marks = int(float(quiz_marks)*25)
                    if quiz_email == mid_email:
                        count+=1
                        gradesheet_ws["M"+str(row)] = quiz_marks
                        infos["C"+str(row-2)] = mid_email
                        infos["D"+str(row-2)] = quiz_bux
                        bot["C"+str(botcount)] = mid_email
                        bot["D"+str(botcount)] = quiz_bux
                        bot["E"+str(botcount)] = quiz_marks
                        break
                break
        if got== False:
            loss_count_sec+=1
            absent.append("ID: "+id+"  Name: "+name)
    print("\n\n\n"+gradesheet_ws.title+":",count_sec,"students out of",loss_count_sec+count_sec,"students attended Mid\n")
    output+="\n\n\n"+gradesheet_ws.title+": "+str(count_sec)+" students out of "+str(loss_count_sec+count_sec)+" students attended Mid.\n\n"
    print("Did Not Attend: \n--------------------------\n\n")
    for abse in absent:
        print(abse)
        output+="\n"+abse
    information.save("information.xlsx")
    gradesheet_workbook.save("new gradesheet.xlsx")
    forbot.save("./mat110_discord_bot/MAT110API.xlsx")
    output+="\n"
print("\nTotal Count:",count)
output+="\n\nTotal Count: "+str(count)
file = open("output.txt","w")
file.write(output)